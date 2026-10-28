"""Build the data files for the Area-based coordination page (abc.html).

Reads the public downloads in raw/ (scripts/abc_fetch.py, run by .github/workflows/abc-data.yml)
and raw/_derived/by-locality.json (scripts/build_abc_points.js), and writes:
  data/abc/areas.json          one row per locality: population, people in need, IDPs, returnees,
                               OCHA 3W presence, basic services, food insecurity, water sources (Aj Jazirah)
  data/abc/ocha-projects.json  Sudan Humanitarian Fund (OCHA) projects 2025 to 2026, placed by the states
                               and localities named in their title and description

Usage: python3 scripts/build_abc_data.py [raw folder]
Needs: openpyxl
"""
import csv
import datetime as dt
import html
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict

import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
RAW = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "raw")
OUT = os.path.join(ROOT, "data", "abc")


def rp(*p):
    return os.path.join(RAW, *p)


def meta(folder):
    with open(rp(folder, "_meta.json"), encoding="utf-8") as f:
        return json.load(f)


def num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "", "-") else None
    except ValueError:
        return None


def rows_csv(path):
    with open(path, encoding="utf-8-sig") as f:
        return [r for r in csv.DictReader(f) if not any(str(v).startswith("#") for v in r.values() if v)]


# ---------- localities on the map ----------
with open(os.path.join(OUT, "localities-geo.json"), encoding="utf-8") as f:
    GEO = json.load(f)
LOCS = [(l["state"], l["name"]) for l in GEO["localities"]]
norm = lambda s: re.sub(r"[^a-z0-9]", "", (s or "").lower())
BY_NAME = defaultdict(list)
for st, nm in LOCS:
    BY_NAME[norm(nm)].append((st, nm))


def find(state, name):
    """Match a (state, locality) from any source to the map's localities."""
    name = re.sub(r"-(GD|NK|SK|SD|ND|WD|ED|CD|WK)$", "", (name or "").strip())
    c = BY_NAME.get(norm(name), [])
    if len(c) == 1:
        return c[0]
    for s, n in c:
        if norm(s) == norm(state):
            return (s, n)
    return None


A = {k: {"n": k[1], "s": k[0]} for k in LOCS}
unmatched = defaultdict(set)

# ---------- 1. Population: UNFPA estimates (COD-PS) ----------
POP_FILE = sorted(f for f in os.listdir(rp("population")) if re.match(r"sdn_admpop_adm2_\d{4}\.csv$", f))[-1]
for r in rows_csv(rp("population", POP_FILE)):
    k = find(r["ADM1_EN"], r["ADM2_EN"])
    if not k:
        unmatched["population"].add(r["ADM2_EN"])
        continue
    a = A[k]
    a.update(p=r["ADM2_PCODE"], ar=r["ADM2_AR"].strip(), pop=round(num(r["T_TL"])), popF=round(num(r["F_TL"])))
    a["popYear"] = int(r["Year"])
PCODE = {a["p"]: k for k, a in A.items() if a.get("p")}

# ---------- 2. People in need: latest HNRP with locality figures ----------
needs_files = sorted(f for f in os.listdir(rp("needs")) if re.match(r"sdn_hpc_needs_api_\d{4}\.csv$", f))
NEEDS_YEAR = None
national = {}
for f in needs_files:
    y = int(re.search(r"(\d{4})", f).group(1))
    rows = rows_csv(rp("needs", f))
    nat = [r for r in rows if not r.get("Admin 2 PCode") and not r.get("Admin 1 PCode") and r.get("Cluster") == "ALL" and (r.get("Description") or "").startswith("Final")]
    if nat:
        national[y] = {"population": num(nat[0]["Population"]), "pin": num(nat[0]["In Need"]), "targeted": num(nat[0]["Targeted"])}
    if any(r.get("Admin 2 PCode") for r in rows):
        NEEDS_YEAR = y
        LOCAL_NEEDS = rows
CLUSTERS = {"FSC": "Food security and livelihoods", "HEA": "Health", "WSH": "Water, sanitation and hygiene", "PRO": "Protection",
            "EDU": "Education", "NUT": "Nutrition", "SHL": "Shelter and non-food items", "ESN": "Shelter and non-food items"}
for r in LOCAL_NEEDS:
    k = PCODE.get(r["Admin 2 PCode"])
    if not k:
        if r["Admin 2 PCode"]:
            unmatched["needs"].add(r["Admin 2 Name"])
        continue
    a = A[k]
    if r["Cluster"] == "ALL" and r["Category"] == "total":
        a["hnoPop"] = round(num(r["Population"]) or 0)
        a["pin"] = round(num(r["In Need"]) or 0)
    elif r["Cluster"] == "ALL" and r["Category"] == "IDP":
        a["pinIdp"] = round(num(r["In Need"]) or 0)
    elif r["Category"] == "total" and r["Cluster"] in CLUSTERS:
        a.setdefault("pinBy", {})[r["Cluster"]] = round(num(r["In Need"]) or 0)
cluster_codes_seen = sorted({r["Cluster"] for r in LOCAL_NEEDS if r["Category"] == "total"})

# ---------- 3. IDPs and returnees: IOM DTM snapshot ----------
dtm_file = [f for f in os.listdir(rp("dtm", "snapshot")) if f.endswith(".xlsx")][0]
wb = openpyxl.load_workbook(rp("dtm", "snapshot", dtm_file), read_only=True, data_only=True)
dtm_dates = {}
for sheet, field in (("IDPs Masterlist (ADMIN2)", "idps"), ("Returnees Masterlist (ADMIN2)", "ret")):
    rows = list(wb[sheet].iter_rows(values_only=True))
    m = re.search(r"(\d{1,2}-[A-Za-z]{3}-\d{4})", str(rows[0][0] or ""))
    dtm_dates[field] = dt.datetime.strptime(m.group(1), "%d-%b-%Y").date().isoformat() if m else None
    for r in rows:
        if not r or not r[3] or not str(r[3]).startswith("SD"):
            continue
        k = PCODE.get(str(r[3]).strip()) or find(r[0], r[2])
        if not k:
            unmatched["dtm"].add(str(r[2]))
            continue
        A[k][field] = round(num(r[4]) or 0)

# ---------- 4. Operational presence: OCHA Sudan 3W, latest three months ----------
w3_file = [f for f in os.listdir(rp("presence", "ocha-sudan-3w")) if f.endswith(".xlsx")][0]
rows = list(openpyxl.load_workbook(rp("presence", "ocha-sudan-3w", w3_file), read_only=True).worksheets[0].iter_rows(values_only=True))[1:]
months = sorted({r[0] for r in rows if r[0]})
LAST3 = months[-3:]
orgs = defaultdict(set)
by_cluster = defaultdict(lambda: defaultdict(set))
by_type = defaultdict(dict)
for mth, st, loc, cl, org, acr, typ in (r[:7] for r in rows):
    if mth not in LAST3 or not loc:
        continue
    k = find(st, loc)
    if not k:
        unmatched["3w"].add(f"{loc} ({st})")
        continue
    oid = (acr or org or "").strip().upper()
    orgs[k].add(oid)
    by_cluster[k][cl].add(oid)
    by_type[k][oid] = typ or "Other"
st_orgs = defaultdict(set)
st_cl = defaultdict(lambda: defaultdict(set))
for k in orgs:
    st_orgs[k[0]] |= orgs[k]
    for c, v in by_cluster[k].items():
        st_cl[k[0]][c] |= v
STATE_W3 = {s: {"w3": len(v), "w3c": {c: len(o) for c, o in sorted(st_cl[s].items(), key=lambda kv: -len(kv[1]))}} for s, v in st_orgs.items()}
for k in orgs:
    A[k]["w3"] = len(orgs[k])
    A[k]["w3c"] = {c: len(v) for c, v in sorted(by_cluster[k].items(), key=lambda kv: -len(kv[1]))}
    A[k]["w3t"] = dict(Counter(by_type[k].values()))

# ---------- 5. Basic services: HeiGIT (facilities and access) and OpenStreetMap ----------
heigit = {r["ADM2_PCODE"]: r for r in rows_csv(rp("heigit", "SDN_ADM2_facilities.csv"))}
access = {r["ADM2_PCODE"]: r for r in rows_csv(rp("heigit", "SDN_ADM2_access.csv"))}
vuln = {r["ADM2_PCODE"]: r for r in rows_csv(rp("heigit", "SDN_ADM2_vulnerability.csv"))}
for pc, k in PCODE.items():
    a = A[k]
    if pc in heigit:
        h = heigit[pc]
        a["hosp"] = int(num(h["hospitals_count"]) or 0)
        a["phc"] = int(num(h["primary_healthcare_count"]) or 0)
        a["edu"] = int(num(h["education_count"]) or 0)
    if pc in access and pc in vuln and num(vuln[pc]["total_pop"]):
        tp = num(vuln[pc]["total_pop"])
        ac = access[pc]
        share = lambda f: round(min(1, (num(ac[f]) or 0) / tp), 3)
        a["accHosp"] = share("access_pop_hospitals_1h")
        a["accPhc"] = share("access_pop_primary_healthcare_1h")
        a["accEdu"] = share("access_pop_education_5km")
        a["hPop"] = round(tp)
derived = json.load(open(rp("_derived", "by-locality.json"), encoding="utf-8"))
for k, a in A.items():
    o = derived["osm"].get(k[1])
    if o and len(BY_NAME[norm(k[1])]) == 1:
        a["osm"] = o

# ---------- 6. Food insecurity: FEWS NET projection (share of the area in each phase) ----------
for k, a in A.items():
    d = derived["ipc"].get(k[1])
    if not d or len(BY_NAME[norm(k[1])]) > 1 or "ML2" not in d:
        continue
    sh = {int(p): v for p, v in d["ML2"].items() if p != "none"}
    if sh:
        a["ipc"] = max(sh, key=lambda p: (sh[p], p))  # phase covering most of the locality
        a["ipcMax"] = max(sh)

# ---------- 7. Water sources: Aj Jazirah state table (August 2026) ----------
with open(os.path.join(OUT, "aj-jazirah-water-sources-2026-08.json"), encoding="utf-8") as f:
    WATER = json.load(f)
types = [t["key"] for t in WATER["types"]]
for row in WATER["localities"]:
    k = find(WATER["state"], row["locality"])
    if k:
        A[k]["water"] = {t: row[t] for t in types if row[t]}
        A[k]["water"]["total"] = sum(row[t] for t in types)

# ---------- 8. Sudan Humanitarian Fund projects (OCHA), from IATI ----------
STATE_ALIASES = {
    "Aj Jazirah": [r"aj[ -]?jazirah?", r"al[ -]?jazirah?", r"gezira", r"al[ -]?gezira", r"gazira", r"el[ -]?gezira"],
    "Blue Nile": [r"blue nile"], "Central Darfur": [r"central darfur"], "East Darfur": [r"east(?:ern)? darfur"],
    "Gedaref": [r"gedaref", r"al[ -]?qadarif", r"gadarif"], "Kassala": [r"kassala"], "Khartoum": [r"khartoum"],
    "North Darfur": [r"north(?:ern)? darfur"], "North Kordofan": [r"north(?:ern)? kordofan"], "Northern": [r"northern state"],
    "Red Sea": [r"red sea"], "River Nile": [r"river nile", r"nahr an nil"], "Sennar": [r"sennar", r"sinnar"],
    "South Darfur": [r"south(?:ern)? darfur"], "South Kordofan": [r"south(?:ern)? kordofan"], "West Darfur": [r"west(?:ern)? darfur"],
    "West Kordofan": [r"west(?:ern)? kordofan"], "White Nile": [r"white nile"],
}
STATE_RX = {s: re.compile(r"\b(?:" + "|".join(v) + r")\b", re.I) for s, v in STATE_ALIASES.items()}
STATE_NAMES = {norm(s) for s in STATE_ALIASES}


def loc_rx(name):
    parts = re.sub(r"^(Al|El|Ad|Ar|As|An|Ash|Az|Ed|Um|Abu)\s+", "", name, flags=re.I)
    prefix = r"(?:(?:al|el|ad|ar|as|an|ash|az|ed)[ -])?" if re.match(r"^(Al|El|Ad|Ar|As|An|Ash|Az|Ed)\s", name, re.I) else ""
    if prefix:
        return re.compile(r"\b" + prefix + re.escape(parts).replace(r"\ ", r"[ -]") + r"\b", re.I)
    return re.compile(r"\b" + re.escape(name).replace(r"\ ", r"[ -]") + r"\b", re.I)


EXTRA_LOC = {("North Darfur", "Al Fasher"): [r"el[ -]?fasher", r"al[ -]?fashir", r"fasher"], ("West Darfur", "Al Geneina"): [r"el[ -]?geneina", r"geneina"],
             ("Central Darfur", "Zalingi"): [r"zalingei", r"zalingi"], ("South Kordofan", "Kadugli"): [r"kadugli"],
             ("South Darfur", "Nyala Janoub"): [r"nyala"], ("North Darfur", "Tawila"): [r"tawila"]}
LOC_RX = []
for st, nm in LOCS:
    if norm(nm) in STATE_NAMES or len(nm) < 5 or nm in ("Abyei PCA area",):
        continue
    rx = [loc_rx(nm)] + [re.compile(r"\b" + p + r"\b", re.I) for p in EXTRA_LOC.get((st, nm), [])]
    LOC_RX.append(((st, nm), rx))

SECTOR_SHORT = {"Water, Sanitation and Hygiene": "WASH", "Emergency Shelter and NFI": "Shelter and NFIs", "Multi-purpose CASH": "Multi-purpose cash",
                "Camp Coordination / Management": "Site management", "Coordination and Support Services": "Coordination"}
projects = []
for f in sorted(os.listdir(rp("iati"))):
    if not f.endswith(".xml"):
        continue
    root = ET.parse(rp("iati", f)).getroot()
    for act in root.findall("iati-activity"):
        t = lambda p: (act.find(p).text or "").strip() if act.find(p) is not None and act.find(p).text else ""
        title = html.unescape(t("title/narrative"))
        desc = re.sub(r"lt/?[a-z]+gt", " ", t("description/narrative"))
        desc = re.sub(r"\s+", " ", html.unescape(desc)).strip()
        text = f"{title}. {desc}"
        states = [s for s, rx in STATE_RX.items() if rx.search(text)]
        locs = [k for k, rxs in LOC_RX if any(rx.search(text) for rx in rxs)]
        # a locality name used in two states (Ar Rahad) counts only where the text also names the state
        locs = [k for k in locs if len(BY_NAME[norm(k[1])]) == 1 or k[0] in states]
        states = sorted(set(states) | {k[0] for k in locs})
        types = sorted({(p.find("narrative").text or "").strip() for p in act.findall("participating-org")
                        if p.get("role") == "4" and p.find("narrative") is not None and (p.find("narrative").text or "").strip()
                        and "Secured" not in (p.find("narrative").text or "")})
        sectors = [SECTOR_SHORT.get(n, n) for n in ((s.find("narrative").text or "").strip() for s in act.findall("sector") if s.get("vocabulary") == "99" and s.find("narrative") is not None) if n]
        dates = {d.get("type"): d.get("iso-date") for d in act.findall("activity-date")}
        budget = sum(num(tr.find("value").text) or 0 for tr in act.findall("transaction") if tr.find("transaction-type") is not None and tr.find("transaction-type").get("code") == "2" and tr.find("value") is not None)
        if not budget:
            budget = sum(num(b.text) or 0 for b in act.findall("budget/value"))
        projects.append({
            "id": t("iati-identifier").replace("XM-OCHA-CBPF-CBPF-", ""), "name": title, "summary": desc[:420] + ("…" if len(desc) > 420 else ""),
            "sectors": sorted(set(sectors)), "orgType": types[0] if len(types) == 1 else ", ".join(types),
            "start": dates.get("2") or dates.get("1"), "end": dates.get("4") or dates.get("3"), "budget": round(budget) if budget else None,
            "states": states, "localities": [{"s": s, "n": n} for s, n in sorted(locs)],
            "nationwide": not states,
        })
projects.sort(key=lambda p: (p["start"] or ""), reverse=True)

# ---------- write ----------
def m_src(folder):
    m = meta(folder)
    return {"title": m["title"], "org": m["org"], "url": m["url"], "updated": (m.get("last_modified") or "")[:10]}


today = dt.date.today().isoformat()
areas = []
for st, nm in LOCS:
    a = A[(st, nm)]
    areas.append({k: v for k, v in a.items() if v not in (None, {}, [])})
out = {
    "updated": today,
    "notes": {
        "population": f"UNFPA population estimates for {next(a['popYear'] for a in areas if a.get('popYear'))}, from the {POP_FILE[-8:-4]} release of the common operational dataset (COD-PS) on HDX.",
        "needs": f"People in need by locality are from the {NEEDS_YEAR} Humanitarian Needs and Response Plan, the latest year with locality figures on HDX. The {max(national)} plan has national figures only so far.",
        "w3": f"Organisations reporting activities to OCHA in {LAST3[0]} to {LAST3[-1]} (3W). Counts, not names.",
        "dtm": f"IOM DTM master lists: IDPs as of {dtm_dates.get('idps')}, returnees as of {dtm_dates.get('ret')}.",
    },
    "national": {str(y): v for y, v in sorted(national.items())},
    "needsYear": NEEDS_YEAR,
    "w3Months": LAST3,
    "dtmDates": dtm_dates,
    "clusters": {c: CLUSTERS[c] for c in cluster_codes_seen if c in CLUSTERS},
    "states": STATE_W3,
    "sources": {
        "population": {**m_src("population"), "file": POP_FILE} if os.path.exists(rp("population", "_meta.json")) else {"title": "Sudan - Subnational Population Statistics", "org": "UNFPA, published by OCHA on HDX", "url": "https://data.humdata.org/dataset/cod-ps-sdn", "file": POP_FILE},
        "needs": m_src("needs") if os.path.exists(rp("needs", "_meta.json")) else {"title": "Sudan: Humanitarian Needs", "org": "OCHA", "url": "https://data.humdata.org/dataset/sudan-humanitarian-needs"},
        "dtm": m_src(os.path.join("dtm", "snapshot")),
        "w3": m_src(os.path.join("presence", "ocha-sudan-3w")),
        "heigit": m_src("heigit"),
        "ipc": m_src("ipc"),
        "osmHealth": m_src(os.path.join("osm", "health")),
        "osmEdu": m_src(os.path.join("osm", "education")),
    },
    "areas": areas,
}
with open(os.path.join(OUT, "areas.json"), "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
with open(os.path.join(OUT, "ocha-projects.json"), "w", encoding="utf-8") as f:
    json.dump({"updated": today, "source": {"title": "Sudan Humanitarian Fund allocations (IATI)", "org": "OCHA, Country-Based Pooled Funds",
               "url": "https://cbpf.data.unocha.org/"}, "note": "States and localities are found by reading each project's title and description, so a few may be missed or misread.",
               "projects": projects}, f, ensure_ascii=False, indent=0)

print("localities", len(areas), "| with population", sum(1 for a in areas if "pop" in a), "| PIN", sum(1 for a in areas if "pin" in a),
      "| IDPs", sum(1 for a in areas if "idps" in a), "| returnees", sum(1 for a in areas if "ret" in a), "| 3W", sum(1 for a in areas if "w3" in a),
      "| HeiGIT", sum(1 for a in areas if "hosp" in a), "| IPC", sum(1 for a in areas if "ipc" in a), "| water", sum(1 for a in areas if "water" in a))
print("SHF projects", len(projects), "| placed", sum(1 for p in projects if p["states"]), "| with localities", sum(1 for p in projects if p["localities"]))
print("needs year", NEEDS_YEAR, "national", national, "3W months", LAST3, "DTM", dtm_dates)
for k, v in unmatched.items():
    print("unmatched", k, sorted(v))
