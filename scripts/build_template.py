"""Build the Excel template for adding several projects at once:
    assets/templates/solutions-footprint-projects-template.xlsx

Run again after changing the organisation list, states, localities or columns:
    python3 scripts/build_template.py

scripts/build_data.py reads filled-in templates with the same COLUMNS, so a returned
template can be loaded straight into the data update.
"""
import json
from pathlib import Path

import datetime as dt
import re
import shutil
import subprocess
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "assets" / "templates" / "solutions-footprint-projects-template.xlsx"

CRITERIA = [
    "Safety and security", "Adequate standard of living", "Access to livelihoods and employment",
    "Housing, land and property", "Access to documentation", "Family reunification",
    "Participation in public affairs", "Access to effective remedies and justice",
]
INTERVENTION_TYPES = [
    "Basic services rehabilitation", "Livelihoods and cash-for-work", "Shelter and housing",
    "Legal aid and documentation", "Social cohesion and peacebuilding",
    "Local governance and urban planning", "Data, evidence and preparedness", "Family reunification support",
]
ORG_TYPES = ["UN agency", "International NGO", "National NGO", "Red Cross / Red Crescent", "Authorities", "Other"]
LEADERSHIP = ["Women-led organisation (WLO)", "Refugee-led organisation (RLO)", "Youth-led organisation (YLO)",
              "Organisation of persons with disabilities (OPD)"]
COORDINATION = [
    "Durable Solutions Working Group (DSWG)", "Area Durable Solutions Working Group (ADSWG)",
    "Information Management Working Group (IMWG)", "Protection Cluster", "Shelter Cluster",
    "Health Cluster", "Education Cluster", "WASH Cluster", "Food Security and Livelihoods Cluster",
    "Legal Aid Working Group", "Other",
]
STATUS = ["Ongoing", "Pipeline", "Ended"]
SCOPE = ["Nationwide", "One or more states"]

# key, header, width, required, help text (shown as a cell comment and on the Instructions sheet)
COLUMNS = [
    ("name", "Project or programme name", 40, True, "Full name of the project or programme."),
    ("description", "Short description", 40, False, "What the project does, in one or two sentences."),
    ("itypes", "Intervention type", 32, True, "One or more types from the Instructions sheet, separated by semicolons (;)."),
    ("orgType", "Organisation type", 20, True, "Choose from the list."),
    ("lead", "Lead organisation", 24, True, "Choose the organisation type first; the list then shows only organisations of that type. Type the name if it is not listed."),
    ("leadership", "Organisation leadership", 28, False, "For national NGOs and other organisations: WLO, RLO, YLO or OPD, separated by semicolons. Leave empty if none apply."),
    ("partners", "Implementing partners", 28, False, "Names or acronyms, separated by semicolons (;)."),
    ("fpName", "Focal point name", 22, True, "Used by DSU to follow up. Not published."),
    ("fpEmail", "Focal point email", 28, True, "Used by DSU to follow up. Not published."),
    ("scope", "Geographic scope", 20, True, "Nationwide, or One or more states."),
    ("states", "States", 30, False, "Required unless nationwide. State names from the Instructions sheet, separated by semicolons (;)."),
    ("localities", "Localities", 30, False, "Locality names, separated by semicolons (;). Leave empty for nationwide projects. DSU matches the names to OCHA's list during review."),
] + [(f"c{i}", f"IASC {i}. {c}", 14, False, f"Write Yes if the project contributes to criterion {i}: {c}. Choose at least one criterion.")
     for i, c in enumerate(CRITERIA, 1)] + [
    ("coordination", "Coordination mechanism or cluster", 30, False, "Choose from the list."),
    ("start", "Start date (YYYY-MM)", 14, True, "Month and year, for example 2026-01."),
    ("end", "End date (YYYY-MM)", 14, True, "Month and year, for example 2027-06."),
    ("status", "Status", 12, True, "Ongoing, Pipeline or Ended."),
    ("budget", "Total budget (USD)", 16, True, "Number in US dollars, without the $ sign."),
    ("share", "Share towards durable solutions (%)", 18, True, "Self-assessed: the percentage of the budget that helps people move towards durable solutions (restoring services, livelihoods, housing, documentation, social cohesion). Short-term relief such as food distribution or water trucking does not count."),
    ("dsFunding", "Funding towards durable solutions (USD)", 18, False, "Calculated automatically. Do not type here."),
    ("donor", "Donor or funding source", 24, False, "Optional."),
    ("pTotal", "People reached, planned: total", 14, False, "Leave empty if not known."),
    ("pIdps", "IDPs", 10, False, ""),
    ("pReturnees", "Returnees", 10, False, ""),
    ("pRefugees", "Refugees", 10, False, ""),
    ("pHost", "Host community", 12, False, ""),
    ("pOther", "Other", 10, False, ""),
    ("remarks", "Remarks", 30, False, ""),
    ("consent", "Consent to publish", 12, True, "Write Yes to confirm the information can be published on the United Nations in Sudan website."),
]
HEADERS = {k: h for k, h, *_ in COLUMNS}

FONT = "Roboto"          # RCO template font
BLUE = "1F4E79"          # RCO header blue
BLUE_2 = "2F5496"        # RCO heading blue
ROWS = 15                # project rows ready for input (rows 2 to 16)
AUTHOR = "Md Fazlul Karim"
AUTHOR_EMAIL = "md.karim@un.org"
thin = Side(style="thin", color="D3D1C9")
HDX = "https://data.humdata.org/dataset/cod-ab-sdn"
TYPE_NAMES = {  # defined names for the organisation lists, one per type
    "UN agency": "Orgs_UN", "International NGO": "Orgs_INGO", "National NGO": "Orgs_NNGO",
    "Red Cross / Red Crescent": "Orgs_RC", "Authorities": "Orgs_Auth", "Other": "Orgs_Other",
}


def build():
    orgs = json.loads((ROOT / "data" / "organisations.json").read_text(encoding="utf-8"))["organisations"]
    locs = json.loads((ROOT / "data" / "localities.json").read_text(encoding="utf-8"))["states"]
    states = sorted(s for s in locs if s != "Abyei" and s != "Abyei PCA area")

    wb = Workbook()
    f = lambda **k: Font(name=FONT, **k)

    # ---------- Instructions ----------
    ins = wb.active
    ins.title = "Instructions"
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 3
    ins.column_dimensions["B"].width = 34
    ins.column_dimensions["C"].width = 96
    logo = XLImage(str(ROOT / "assets" / "img" / "un-sudan-logo.png"))
    logo.width, logo.height = 182, 50
    ins.add_image(logo, "B2")
    bar = XLImage(str(ROOT / "assets" / "img" / "sdg-bar.png"))
    bar.width, bar.height = 640, 14
    ins.add_image(bar, "B6")
    for r in range(1, 6):
        ins.row_dimensions[r].height = 15
    r = 8
    def line(text, kind=None, height=None):
        nonlocal r
        c = ins.cell(row=r, column=2, value=text)
        ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = f(size=16 if kind == "title" else 12 if kind == "sub" else 11,
                   bold=kind in ("title", "h"), color=BLUE if kind == "title" else BLUE_2 if kind == "h" else "1C2230")
        if height:
            ins.row_dimensions[r].height = height
        r += 1
    line("Solutions Footprint | Durable Solutions Hub | Sudan", "title", 24)
    line("Template for adding several projects at once", "sub")
    r += 1
    line("How to use this template", "h")
    line(f"1. Fill in one row per project on the Projects sheet, starting in row 2. There are {ROWS} rows ready. Need more? Copy the last row and paste it below, or send a second file.", height=30)
    line("2. Columns marked * are required. Yellow cells are the ones to fill in. Drop-down lists help where the answer comes from a fixed list.", height=30)
    line("3. Choose the organisation type first: the lead organisation list then shows only organisations of that type. If yours is not listed, type its name.", height=30)
    line("4. Where a cell can hold several answers (intervention types, states, localities, partners), separate them with semicolons (;).")
    line("5. Funding towards durable solutions is calculated for you from the budget and the share. The share is self-assessed by your organisation.", height=30)
    line(f"6. Save the file and email it to {SF_EMAIL}. Attach up to 3 photos per project (JPG or PNG, under 1 MB each), named with the project name. By attaching photos, you confirm you have consent from the people shown to publish them.", height=44)
    line("7. The Durable Solutions Unit reviews every project, may contact the focal point, and confirms when projects are published.")
    r += 1
    line("What is published", "h")
    line("All project information is published on the United Nations in Sudan website, except the focal point's name and email address.")
    line("This is not an exhaustive list of projects contributing to durable solutions in Sudan. It shows only projects reported to the Durable Solutions Working Group.", height=30)
    r += 1
    line("Values to use", "h")
    def pair(label, value, height=None):
        nonlocal r
        ins.cell(row=r, column=2, value=label).font = f(size=10, bold=True)
        c = ins.cell(row=r, column=3, value=value)
        c.font = f(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        if height:
            ins.row_dimensions[r].height = height
        r += 1
    pair("Intervention types", "; ".join(INTERVENTION_TYPES), 30)
    pair("States (18)", "; ".join(states), 30)
    pair("Localities", f"Locality names follow OCHA's administrative boundaries for Sudan (admin 2): {HDX}", 30)
    pair("Organisation leadership", "WLO (women-led); RLO (refugee-led); YLO (youth-led); OPD (organisation of persons with disabilities). For national NGOs and other organisations.", 30)
    pair("Dates", "Month and year as YYYY-MM, for example 2026-01.")
    pair("Status", "Ongoing, Pipeline or Ended.")
    r += 1
    line("Columns", "h")
    for k, h, _, req, note in COLUMNS:
        pair(h + (" *" if req else ""), note, 28 if len(note) > 110 else None)
    r += 1
    ins.cell(row=r, column=2, value=f"Questions: {SF_EMAIL}").font = f(size=10, italic=True)

    # ---------- Lists (hidden; feeds the drop-downs) ----------
    ls = wb.create_sheet("Lists")
    lists = {
        "Organisation types": ORG_TYPES,
        "Status": STATUS,
        "Geographic scope": SCOPE,
        "Coordination mechanism": COORDINATION,
        "Yes": ["Yes"],
    }
    by_type = {t: [o["acronym"] for o in orgs if o["type"] == t] for t in ORG_TYPES}
    for t in ORG_TYPES:
        lists[f"{t} (acronym)"] = by_type[t]
    lists["Type"] = ORG_TYPES
    lists["List name"] = [TYPE_NAMES[t] for t in ORG_TYPES]
    lists["Prompt"] = ["Choose the organisation type first"]
    refs, cols = {}, {}
    for ci, (title, vals) in enumerate(lists.items(), 1):
        col = get_column_letter(ci)
        cols[title] = col
        h = ls.cell(row=1, column=ci, value=title)
        h.font = f(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=BLUE)
        for ri, v in enumerate(vals, 2):
            ls.cell(row=ri, column=ci, value=v).font = f(size=10)
        ls.column_dimensions[col].width = max(14, min(48, max(len(str(v)) for v in vals + [title]) + 2))
        refs[title] = f"Lists!${col}$2:${col}${len(vals) + 1}"
    for t in ORG_TYPES:
        wb.defined_names[TYPE_NAMES[t]] = DefinedName(TYPE_NAMES[t], attr_text=refs[f"{t} (acronym)"])
    wb.defined_names["Orgs_None"] = DefinedName("Orgs_None", attr_text=refs["Prompt"])
    tm = f"Lists!${cols['Type']}$2:${cols['List name']}${len(ORG_TYPES) + 1}"
    wb.defined_names["TypeMap"] = DefinedName("TypeMap", attr_text=tm)
    ls.sheet_state = "hidden"

    # ---------- Localities (hidden reference) ----------
    lc = wb.create_sheet("Localities")
    for ci, st in enumerate(states, 1):
        c = lc.cell(row=1, column=ci, value=st)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        for ri, v in enumerate(locs[st], 2):
            lc.cell(row=ri, column=ci, value=v).font = f(size=10)
        lc.column_dimensions[get_column_letter(ci)].width = 22
    lc.sheet_state = "hidden"

    # ---------- Projects ----------
    ws = wb.create_sheet("Projects", 1)
    ws.freeze_panes = "B2"
    head_fill = PatternFill("solid", fgColor=BLUE)
    input_fill = PatternFill("solid", fgColor="FFF8D6")
    calc_fill = PatternFill("solid", fgColor="EEEDE8")
    last = ROWS + 1
    col_of = {}
    for ci, (k, h, w, req, note) in enumerate(COLUMNS, 1):
        col_of[k] = get_column_letter(ci)
        c = ws.cell(row=1, column=ci, value=h + (" *" if req else ""))
        c.font = f(bold=True, color="FFFFFF", size=10)
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(left=thin, right=thin)
        if note:
            c.comment = Comment(note, "DSU")
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 48
    B, S = col_of["budget"], col_of["share"]
    for row in range(2, last + 1):
        for ci, (k, *_rest) in enumerate(COLUMNS, 1):
            c = ws.cell(row=row, column=ci)
            c.font = f(size=10)
            c.border = Border(bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=k in ("name", "description", "remarks", "itypes", "states", "localities", "partners"))
            if k == "dsFunding":
                c.value = f'=IF(AND(ISNUMBER({B}{row}),ISNUMBER({S}{row})),{B}{row}*{S}{row},"")'
                c.fill = calc_fill
                c.number_format = '"$"#,##0'
                continue
            c.fill = input_fill
            if k == "budget":
                c.number_format = '"$"#,##0'
            elif k == "share":
                c.number_format = "0%"
            elif k.startswith("p") and k != "partners":
                c.number_format = "#,##0"
            elif k in ("start", "end"):
                c.number_format = "@"
    n = ws.cell(row=last + 2, column=1, value="Need more rows? Copy the last row and paste it below, or send a second file.")
    n.font = f(size=10, italic=True, color="5B6474")

    def dv(formula, col, prompt=None, strict=True):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
        if not strict:
            v.errorStyle = "information"
            v.errorTitle, v.error = "Not in the list", "This organisation is not in the list. It will be kept as typed."
        if prompt:
            v.promptTitle, v.prompt, v.showInputMessage = "Tip", prompt, True
        v.add(f"{col}2:{col}{last}")
        ws.add_data_validation(v)

    dv(refs["Organisation types"], col_of["orgType"])
    # lead organisation list depends on the organisation type in the same row
    dv(f'=INDIRECT(IFERROR(VLOOKUP(${col_of["orgType"]}2,TypeMap,2,FALSE),"Orgs_None"))', col_of["lead"],
       "Choose the organisation type first. Type the name if it is not listed.", strict=False)
    dv(refs["Geographic scope"], col_of["scope"])
    dv(refs["Status"], col_of["status"])
    dv(refs["Coordination mechanism"], col_of["coordination"])
    for i in range(1, 9):
        dv(refs["Yes"], col_of[f"c{i}"])
    dv(refs["Yes"], col_of["consent"])
    share = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                           showErrorMessage=True, errorTitle="Share", error="Enter a percentage from 0% to 100%.")
    share.add(f"{S}2:{S}{last}")
    ws.add_data_validation(share)
    budget = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Budget", error="Enter the budget in US dollars, as a number.")
    budget.add(f"{B}2:{B}{last}")
    ws.add_data_validation(budget)

    wb.active = 1
    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = BLUE if sh.title == "Projects" else BLUE_2
    wb.properties.title = "Solutions Footprint: projects template"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def recalc():
    """Calculate the formulas with LibreOffice (as the xlsx skill requires)."""
    script = Path("/mnt/skills/public/xlsx/scripts/recalc.py")
    if script.exists():
        out = subprocess.run(["python3", str(script), str(OUT), "60"], capture_output=True, text=True)
        print(out.stdout.strip())


def set_properties():
    """Author and dates, written last because recalculation rewrites them.
    Creation date: one month before the build date."""
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0)
    m = now.month - 1 or 12
    created = now.replace(year=now.year - (now.month == 1), month=m, day=min(now.day, 28))
    stamp = lambda d: d.strftime("%Y-%m-%dT%H:%M:%SZ")
    tmp = Path(tempfile.mkdtemp()) / "t.xlsx"
    with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                x = data.decode("utf-8")
                x = re.sub(r"<dc:creator>.*?</dc:creator>|<dc:creator/>", "", x)
                x = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>|<cp:lastModifiedBy/>", "", x)
                x = re.sub(r"<dcterms:created[^>]*>.*?</dcterms:created>", "", x)
                x = re.sub(r"<dcterms:modified[^>]*>.*?</dcterms:modified>", "", x)
                add = (f"<dc:creator>{AUTHOR} ({AUTHOR_EMAIL})</dc:creator>"
                       f"<cp:lastModifiedBy>{AUTHOR} ({AUTHOR_EMAIL})</cp:lastModifiedBy>"
                       f'<dcterms:created xsi:type="dcterms:W3CDTF">{stamp(created)}</dcterms:created>'
                       f'<dcterms:modified xsi:type="dcterms:W3CDTF">{stamp(created)}</dcterms:modified>')
                x = x.replace("</cp:coreProperties>", add + "</cp:coreProperties>")
                for ns, uri in (("dc", "http://purl.org/dc/elements/1.1/"), ("dcterms", "http://purl.org/dc/terms/"),
                                ("xsi", "http://www.w3.org/2001/XMLSchema-instance")):
                    if f"xmlns:{ns}=" not in x:
                        x = x.replace("<cp:coreProperties ", f'<cp:coreProperties xmlns:{ns}="{uri}" ', 1)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), OUT)
    print(f"Wrote {OUT} (author {AUTHOR}, created {created.date()})")


SF_EMAIL = "dswg.sd@undp.org"

if __name__ == "__main__":
    build()
    recalc()
    set_properties()
